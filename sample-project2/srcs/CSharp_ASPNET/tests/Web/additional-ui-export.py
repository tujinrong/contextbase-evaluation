from pathlib import Path
import asyncio,json,sys,importlib.util,subprocess,uuid,time
from playwright.async_api import async_playwright,expect
import additional_runtime as rt
s=importlib.util.spec_from_file_location('cases',Path(__file__).with_name('additional-ui-cases.py'));h=importlib.util.module_from_spec(s);s.loader.exec_module(h)
ROOT=rt.ROOT;OUT=h.OUT;SOURCES=[Path(__file__),Path(__file__).with_name('additional-ui-cases.py'),Path(__file__).with_name('additional_runtime.py'),ROOT/'tests/RuntimeObserver/Program.cs']
async def selected(p):
    await h.search_page(p);await p.locator('#f03-i01').fill('000001');await h.search(p)
    async with p.expect_response(lambda r:r.url.endswith('/employees/select')):await p.locator('#search-rows input[type=radio]').check()
    await expect(p.locator('#employee-export')).to_be_enabled();return await h.vm(p)
class RowLock:
    def __init__(self):
        self.token=uuid.uuid4().hex
        query="SET CONTEXT_INFO 0x"+self.token+";\nGO\nSET NOCOUNT ON;IF DB_NAME()<>N'"+rt.DB+"' THROW 51000,'Wrong own DB',1;BEGIN TRAN;SELECT employee_id FROM dbo.m_employee WITH(TABLOCKX,HOLDLOCK) WHERE employee_no='000001';WAITFOR DELAY '00:00:06';ROLLBACK;"
        self.process=subprocess.Popen(['sqlcmd','-S',r'.\SQLEXPRESS','-E','-C','-b','-d',rt.DB,'-Q',query],stdout=subprocess.DEVNULL,stderr=subprocess.PIPE)
    async def ready(self):
        for _ in range(50):
            found=rt.rows("SELECT COUNT(*) AS n,MAX(s.session_id) AS own_session_id FROM sys.dm_tran_locks l JOIN sys.dm_exec_sessions s ON s.session_id=l.request_session_id WHERE l.resource_database_id=DB_ID() AND l.request_mode='X' AND SUBSTRING(s.context_info,1,16)=0x"+self.token)[0]
            if found['n']:self.session=found['own_session_id'];return found['n']
            await asyncio.sleep(.05)
        raise AssertionError('Own row lock not observed')
    async def reader(self):
        for _ in range(40):
            value=rt.sql("SELECT r.session_id,r.blocking_session_id,r.wait_type,t.text AS actual_sql FROM sys.dm_exec_requests r CROSS APPLY sys.dm_exec_sql_text(r.sql_handle)t WHERE r.database_id=DB_ID() AND r.blocking_session_id="+str(self.session)+" FOR JSON PATH;")
            if value:
                rows=json.loads(''.join(value.splitlines()));assert rows and all(x['blocking_session_id']==self.session for x in rows);return rows
            await asyncio.sleep(.03)
        raise AssertionError('Report read was not observed behind the own row lock')
    async def finish(self):
        while self.process.poll() is None:await asyncio.sleep(.05)
        assert self.process.returncode==0,self.process.stderr.read().decode(errors='replace')
def db():return {'employee':rt.rows("SELECT * FROM dbo.m_employee WHERE employee_no='000001'"),'careers':rt.rows("SELECT c.* FROM dbo.t_career c JOIN dbo.m_employee e ON e.employee_id=c.employee_id WHERE e.employee_no='000001' ORDER BY c.career_id")}
async def download_duplicate(p,c,requests,dialogs):
    before=await selected(p);raw_before=await h.raw(p);saved=db();lock=RowLock();start=len(requests)
    try:
        lock_count=await lock.ready()
        async with p.expect_request(lambda r:r.url.endswith('/employees/export')) as request_wait:
            async with p.expect_download() as download_wait:
                await p.locator('#employee-export').click();request=await request_wait.value;payload=request.post_data_json
                blocked_reader=await lock.reader();current=h.observer.snapshot(before['stateId'])
                disabled=await p.locator('#employee-export').is_disabled()
                assert current['fields']['Revision']>before['stateRevision'] and disabled,{'before':before['stateRevision'],'server':current['fields']['Revision'],'disabled':disabled,'payload':payload,'vm':await h.vm(p),'message':await p.locator('#message').inner_text()}
                await p.locator('#employee-export').click(force=True);await asyncio.sleep(.1);ui_requests=sum(x['path']=='/employees/export' for x in requests[start:]);assert ui_requests==1
                duplicate={**payload,'stateRevision':current['fields']['Revision']};rejected=await h.post(p,'/employees/export',duplicate);assert rejected['status']==503 and rejected['body']['outcome']['messageId']=='MSG-029'
                await lock.finish();download=await download_wait.value
        file=OUT/'selected-saved-employee.xlsx';await download.save_as(file);assert await download.failure() is None
    finally:await lock.finish()
    import openpyxl
    workbook=openpyxl.load_workbook(file,read_only=True,data_only=False);sheet_count=len(workbook.sheetnames);values=[str(cell.value or '') for sheet in workbook for row in sheet.iter_rows() for cell in row];workbook.close();current=await h.vm(p);notice=await p.locator('#message').inner_text();included=all(any(row['project_name'] in value for value in values) for row in saved['careers']);assert sheet_count==2 and included and saved==db();raw_same=await h.raw(p)==raw_before and current['selectedEmployeeId']==before['selectedEmployeeId'];assert raw_same and '開始' in notice
    props={'savedSelectedEmployees':len(before['rows']),'sheetCount':sheet_count,'allSavedCareersIncluded':included,'searchRawAndSelectionRetained':raw_same,'downloadStartClaimedClientSaveComplete':'保存完了' in notice or '保存しました' in notice,'duplicateGenerationRejected':ui_requests==1 and rejected['status']==503};assert props['savedSelectedEmployees']==1 and not props['downloadStartClaimedClientSaveComplete']
    manual_start=len(requests)
    async with p.expect_download() as manual_wait:await p.locator('#employee-export').click()
    manual=await manual_wait.value;assert await manual.failure() is None;await expect(p.locator('#employee-export')).to_be_enabled();manual_vm=await h.vm(p);manual_server=h.observer.snapshot(before['stateId']);await asyncio.sleep(.7);manual_count=sum(x['path']=='/employees/export' for x in requests[manual_start:]);assert manual_count==1 and manual_vm['ui']['stateRevision']==manual_server['fields']['Revision'] and saved==db()
    return {'adopt':{'DUI12.export':props},'selectedBefore':before,'rawAfter':await h.raw(p),'selectedAfter':current,'savedDb':saved,'notice':notice,'outboundUiGenerationRequestsBeforeProtocolProbe':ui_requests,'busyDuplicateResponse':rejected,'ownRowLockXCount':lock_count,'actualBlockedReportReader':blocked_reader,'workbook':h.proof(file),'postDuplicateManualExportCount':manual_count,'postDuplicateClientRevision':manual_vm['ui']['stateRevision'],'postDuplicateServerRevision':manual_server['fields']['Revision'],'automaticRetryObservationWindowMs':700,'scope':'Actual selected-person download/readback, DB careers membership, UI duplicate pointer rejection plus concurrent server503/MSG029 while a read-only own table lock delays generation, then another manual export succeeds after registry release. Observer acquires the state Gate only after the real SQL blocked reader proves command acceptance. No Office Automation.'}

async def cancel_retry(p,c,requests,dialogs):
    state=await selected(p);saved=db();lock=RowLock();start=len(requests)
    try:
        locks=await lock.ready()
        async with p.expect_request(lambda r:r.url.endswith('/employees/export')) as request_wait:await p.locator('#employee-export').click()
        payload=(await request_wait.value).post_data_json;await expect(p.locator('#export-cancel')).to_be_visible()
        blocked_reader=await lock.reader();server=h.observer.snapshot(state['stateId'])
        duplicate=await h.post(p,'/employees/export',{**payload,'stateRevision':server['fields']['Revision']});assert duplicate['status']==503 and duplicate['body']['outcome']['messageId']=='MSG-029'
        async with p.expect_response(lambda r:r.url.endswith('/employees/export/cancel')) as canceled_wait:await p.locator('#export-cancel').click()
        canceled=await canceled_wait.value;body=await canceled.json();assert canceled.status==200 and body['phase']=='CancelledBeforeStart';await expect(p.locator('#export-cancel')).not_to_be_visible();await lock.finish();after=await h.vm(p);server_after=h.observer.snapshot(state['stateId']);assert db()==saved
        manual_start=len(requests)
        async with p.expect_download() as download_wait:
            async with p.expect_response(lambda r:r.url.endswith('/employees/export')) as retry_wait:await p.locator('#employee-export').click()
            retry=await retry_wait.value;retry_type=retry.headers.get('content-type');retry_body=await retry.json() if retry_type and retry_type.startswith('application/json') else None
            download=await download_wait.value;assert await download.failure() is None
        await expect(p.locator('#employee-export')).to_be_enabled();await asyncio.sleep(.7);manual_count=sum(x['path']=='/employees/export' for x in requests[manual_start:]);assert manual_count==1
        details={'cancelResponse':body,'clientRevisionAfterCancel':after['ui']['stateRevision'],'serverRevisionAfterCancel':server_after['fields']['Revision'],'busyDuplicateResponse':duplicate,'retryHttpStatus':retry.status,'retryContentType':retry_type,'retryBody':retry_body,'manualExportPosts':manual_count,'automaticRetryObservationWindowMs':700,'dbUnchanged':db()==saved,'ownRowLockXCount':locks,'actualBlockedReportReader':blocked_reader,'scope':'Actual Reading, concurrent duplicate503/MSG029, cancel then immediate UI retry; same-state UI revision, complete download and saved DB readback.'}
        assert retry.status==200 and retry_type.startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet') and after['ui']['stateRevision']==server_after['fields']['Revision'],details
        return details
    finally:await lock.finish()
async def failure_retry(p,c,requests,dialogs):
    state=await selected(p);before_raw=await h.raw(p);saved=db();at=len(requests)
    rt.sql("EXEC sys.sp_rename N'dbo.t_career',N't_career_EXPORT_fault',N'OBJECT';")
    try:
        async with p.expect_response(lambda r:r.url.endswith('/employees/export')) as failed_wait:await p.locator('#employee-export').click()
        failed=await failed_wait.value;body=await failed.json();await expect(p.locator('#employee-export')).to_be_enabled();after=await h.vm(p);server=h.observer.snapshot(state['stateId']);assert not failed.ok and body['outcome']['messageId']=='MSG-028' and after['ui']['stateRevision']==server['fields']['Revision'] and await h.raw(p)==before_raw
    finally:rt.sql("EXEC sys.sp_rename N'dbo.t_career_EXPORT_fault',N't_career',N'OBJECT';")
    async with p.expect_download() as download_wait:
        async with p.expect_response(lambda r:r.url.endswith('/employees/export')) as retry_wait:await p.locator('#employee-export').click()
        retry=await retry_wait.value;assert retry.status==200
        download=await download_wait.value;assert await download.failure() is None
    assert db()==saved
    return {'failureStatus':failed.status,'failureBody':body,'clientAfterFailure':after,'serverAfterFailure':server,'rawBefore':before_raw,'retryStatus':retry.status,'actualExportPosts':sum(x['path']=='/employees/export' for x in requests[at:]),'dbUnchanged':db()==saved,'scope':'Actual owned missing child table triggers report-read failure; restored in finally before same-screen manual retry. No automatic retry or DB data write.'}
async def main():
    start=rt.artifact();anchors={'contractSha256':rt.sha(ROOT/'tests/ui-contract.json'),'inputManifestSha256':rt.sha(ROOT/'build/generation/inputs.json')};assert start['sha256']==h.EXPECTED_ARTIFACT;hashes=[h.proof(p) for p in SOURCES];runner=rt.sha(ROOT/'.tmp/bin/RuntimeObserver/Release/net10.0/RuntimeObserver.dll');h.observer=rt.Observer()
    try:
        async with async_playwright() as pw:
            b=await pw.chromium.launch();selected_cases=sys.argv[1:]
            if not selected_cases or 'download' in selected_cases:await h.case(b,'export-download-duplicate',download_duplicate)
            if not selected_cases or 'cancel' in selected_cases:await h.case(b,'export-cancel-retry',cancel_retry)
            if not selected_cases or 'failure' in selected_cases:await h.case(b,'export-failure-retry',failure_retry)
            await b.close()
    finally:stop=h.observer.stop()
    assert anchors=={'contractSha256':rt.sha(ROOT/'tests/ui-contract.json'),'inputManifestSha256':rt.sha(ROOT/'build/generation/inputs.json')};assert start==rt.artifact() and hashes==[h.proof(p) for p in SOURCES] and runner==rt.sha(ROOT/'.tmp/bin/RuntimeObserver/Release/net10.0/RuntimeObserver.dll')
    record={'runId':'sp2-ui-regenerate-20261029','artifact':start,'contractSha256':rt.sha(ROOT/'tests/ui-contract.json'),'inputManifestSha256':rt.sha(ROOT/'build/generation/inputs.json'),'anchorsAtStart':anchors,'anchorsAtEnd':{'contractSha256':rt.sha(ROOT/'tests/ui-contract.json'),'inputManifestSha256':rt.sha(ROOT/'build/generation/inputs.json')},'measurementSources':hashes,'observerRuntime':{'sha256':runner,'ownership':h.observer.ownership,'stop':stop},'results':h.results,'observations':list(h.observations.values()),'pageErrors':h.errors,'formalScore':None};(OUT/'export.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf8');print(json.dumps({'pass':sum(x['status']=='PASS' for x in h.results),'fail':sum(x['status']=='FAIL' for x in h.results),'errors':h.errors}));return int(any(x['status']=='FAIL' for x in h.results) or bool(h.errors))
if __name__=='__main__':sys.exit(asyncio.run(main()))
