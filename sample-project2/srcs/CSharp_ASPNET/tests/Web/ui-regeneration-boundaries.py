from pathlib import Path
import asyncio,json,hashlib,datetime,sys,importlib.util,traceback
from playwright.async_api import async_playwright,expect
import fixture_sql as db

ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'.tmp/build-agent/boundaries';OUT.mkdir(parents=True,exist_ok=True)
s=importlib.util.spec_from_file_location('h',Path(__file__).with_name('ui-regeneration-history.py'));h=importlib.util.module_from_spec(s);s.loader.exec_module(h);h.h.OUT=OUT
results=[];observed={};errors=[];marker='UIBOUNDARYR20261029'
async def record(name,page,fn):
    try:
        detail=await fn();results.append({'id':name,'status':'PASS','detail':detail});print('PASS '+name,flush=True)
    except Exception as e:
        image=await h.h.shot(page,'failure-'+name);results.append({'id':name,'status':'FAIL','error':str(e),'traceback':traceback.format_exc(),'image':image});print('FAIL '+name+' '+json.dumps(traceback.format_exc(),ensure_ascii=True),flush=True)
async def criteria(p,name):
    await h.h.type_text(p,'#f03-i01','');await h.h.type_text(p,'#f03-i02',name);await h.search(p)
def add(n):
    assert n in (200,201)
    db.execute("IF EXISTS(SELECT 1 FROM dbo.m_employee WHERE employee_no BETWEEN '830001' AND '830201') THROW 51000,'Boundary numbers already occupied',1;"+";WITH N AS(SELECT TOP("+str(n)+") ROW_NUMBER() OVER(ORDER BY object_id) AS n FROM sys.all_objects) INSERT dbo.m_employee(employee_no,employee_name,employee_kana,department_code,employment_status,join_date,updated_at,updated_by) SELECT RIGHT('000000'+CONVERT(varchar(6),830000+n),6),N'"+marker+"-'+CONVERT(nvarchar(3),n),N'テスト','01','01','20200401',SYSUTCDATETIME(),N'SP2STAFF' FROM N;")
def cleanup():
    # Only the exact new test-number range carrying this run's marker; no user rows.
    db.execute("IF EXISTS(SELECT 1 FROM dbo.t_career c JOIN dbo.m_employee e ON e.employee_id=c.employee_id WHERE e.employee_no BETWEEN '830001' AND '830201') THROW 51000,'Unexpected boundary career; refusing cleanup',1; DELETE dbo.m_employee WHERE employee_no BETWEEN '830001' AND '830201' AND employee_name LIKE N'"+marker+"%';")
async def zero(p):
    await criteria(p,marker+'NONE');m=await h.model(p);assert m['resultStatus']=='Success' and m['totalCount']==0 and len(m['rows'])==0 and m['selectedEmployeeId'] is None
    assert await p.locator('#employee-export').is_disabled();assert '0' in await p.locator('#search-count').inner_text()
    observed['DUI06.zero']={'resultStatus':m['resultStatus'],'totalCount':m['totalCount'],'rows':len(m['rows']),'selected':False,'exportEnabled':False,'unsearchedConflated':False};return observed['DUI06.zero']
async def two_hundred(p):
    await criteria(p,marker);m=await h.model(p);nos=await p.locator('#search-rows tr td:nth-child(2)').all_text_contents();assert m['totalCount']==200 and len(nos)==200 and nos==sorted(nos) and len(set(nos))==200
    box=p.locator('[data-item-id="F03-I08"]');assert await box.evaluate('(e)=>getComputedStyle(e).overflowY') in ('auto','scroll');assert await box.evaluate('(e)=>e.scrollHeight>e.clientHeight')
    last=p.locator('#search-rows tr').last;await last.locator('input[type=radio]').check();await expect(p.locator('#employee-export')).to_be_enabled();await last.locator('a').click();await h.h.ready(p);assert '830200' in await p.locator('[data-ui="employee-identity"]').inner_text()
    await h.h.go(p,'#employee-cancel',r'/employees(?:\?|$)');m=await h.model(p);assert m['totalCount']==200 and m['selectedEmployeeId'] is not None;assert await box.evaluate('(e)=>e.scrollTop>0')
    observed['DUI06.200']={'displayedRows':len(nos),'stableSort':'employeeNo ASC','pagingControls':await p.locator('[data-page],.pagination').count(),'internalVerticalScroll':True,'range':'all fetched rows','selectionPreserved':True};return observed['DUI06.200']
async def above(p,context):
    db.execute("INSERT dbo.m_employee(employee_no,employee_name,employee_kana,department_code,employment_status,join_date,updated_at,updated_by) VALUES('830201',N'"+marker+"-201',N'テスト','01','01','20200401',SYSUTCDATETIME(),N'SP2STAFF');")
    await h.search(p);m=await h.model(p);assert m['resultStatus']=='LimitExceeded' and not m['rows'] and m['totalCount'] is None and m['selectedEmployeeId'] is None;assert await p.locator('#employee-export').is_disabled()
    await expect(p.locator('#message')).to_contain_text('200');before=await h.raw(p);warm,stats=await h.warm(context);await p.reload();await h.h.ready(p);after=await h.model(p);afterstats=db.search_counter();counter=h.delta(stats,afterstats);await warm.close();assert counter==0,'SQL zero unavailable or unexpectedly executed'
    assert after['resultStatus']=='LimitExceeded' and await h.raw(p)==before
    observed['DUI06.201']={'resultStatus':m['resultStatus'],'rows':0,'oldSelectionCleared':True,'exportEnabled':False,'criteriaRetained':True,'maximum200ExplanationVisible':True,'partial200Conflated':False}
    original=m['ui'];probe=await h.post(p,'/ui/attach',{'stateId':original['stateId'],'tabId':original['tabId'],'documentId':original['documentId'],'documentLease':original['documentLease']});assert probe['status']==200
    scroll=await p.locator('[data-item-id="F03-I08"]').evaluate('(e)=>({top:e.scrollTop,max:Math.max(0,e.scrollHeight-e.clientHeight)})');clamped=scroll['top']==min(m['scrollTop'],scroll['max']);assert clamped
    observed['DUI11.history.LimitExceeded']={'restoredStatus':after['resultStatus'],'rawFourValuesPreserved':await h.raw(p)==before,'appliedSearchRequests':counter,'newStateAllocated':after['stateId']!=m['stateId'],'originalLiveStateUnchanged':probe['status']==200,'scrollRetainedOrClamped':clamped}
    return {'limit':observed['DUI06.201'],'history':observed['DUI11.history.LimitExceeded'],'counterBefore':stats,'counterAfter':afterstats}
async def keyboard(p):
    await p.goto(h.h.runtime['url']+'/employees');await h.h.ready(p);await p.locator('#search-clear').click();await expect(p.locator('#search-clear')).to_be_enabled();await p.locator('#employee-new').focus();sequence=[]
    for _ in range(7):
        sequence.append(await p.evaluate('document.activeElement.id'));await p.keyboard.press('Tab')
    assert sequence==['employee-new','f03-i01','f03-i02','f03-i03','f03-i04','search-run','search-clear'],sequence
    await p.locator('#f03-i01').focus();await p.keyboard.type('000001');await p.keyboard.press('Enter');await p.wait_for_selector('#search-rows tr');await h.h.go(p,'#search-rows a',r'/employees/[0-9]+/edit(?:\?|$)')
    tabs=p.locator('[role=tab]');await tabs.first.focus();await p.keyboard.press('ArrowRight');await expect(tabs.last).to_be_focused();await p.keyboard.press('Enter');await expect(tabs.last).to_have_attribute('aria-selected','true');await p.keyboard.press('Home');await expect(tabs.first).to_be_focused();await p.keyboard.press('End');await expect(tabs.last).to_be_focused()
    await tabs.first.click();await expect(tabs.first).to_have_attribute('aria-selected','true');await expect(p.locator('#f04-i08')).to_be_enabled();await p.locator('#f04-i08').focus();await p.keyboard.press('End');await p.keyboard.press('Enter');await p.keyboard.type('改行検査');v=await p.locator('#f04-i08').input_value();assert '\n改行検査' in v
    await p.locator('#employee-cancel').click();await expect(p.locator('dialog')).to_be_visible();await p.keyboard.press('Escape');await expect(p.locator('dialog')).not_to_be_visible();assert await p.locator('#f04-i08').input_value()==v
    return {'naturalF03TabPrefix':' > '.join(sequence),'searchEnterPerformed':True,'nativeArrowHomeEndTabBehavior':True,'textareaEnterIsNewline':True,'escapePreservesInput':True,'physicalImeClipboard':'NOT_EXECUTED'}
async def download(p):
    await p.goto(h.h.runtime['url']+'/employees');await h.h.ready(p);await p.locator('#search-clear').click();await expect(p.locator('#search-clear')).to_be_enabled();await h.h.type_text(p,'#f03-i01','000001');await h.search(p);await p.locator('#search-rows input[type=radio]').first.check();await expect(p.locator('#employee-export')).to_be_enabled()
    async with p.expect_download() as pending:await p.locator('#employee-export').click()
    d=await pending.value;file=OUT/'selected-employee.xlsx';await d.save_as(file);assert await d.failure() is None
    import openpyxl
    w=openpyxl.load_workbook(file,read_only=True,data_only=False);sheets=w.sheetnames;assert len(sheets)==2;values=[[str(c.value or '') for row in s.iter_rows() for c in row] for s in w];assert any('000001' in v for rows in values for v in rows);w.close()
    return {'suggestedFilename':d.suggested_filename,'sheetCount':len(sheets),'selectedEmployeeNoPresent':True,'file':str(file),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'scope':'Headless Chromium download plus saved XLSX readback; no OS Save dialog or physical print claim.'}
async def main():
    artifact=h.func.artifact();initial=db.rows('SELECT COUNT(*) AS n FROM dbo.m_employee')[0]['n'];add(200)
    try:
        async with async_playwright() as pw:
            browser=await pw.chromium.launch();context=await browser.new_context(ignore_https_errors=True,viewport={'width':1366,'height':768},accept_downloads=True);p=await context.new_page();p.on('pageerror',lambda e:errors.append(str(e)));await h.h.login(p,'staff');await h.prepare(p)
            await record('ZERO-RESULT',p,lambda:zero(p));await record('EXACT-200',p,lambda:two_hundred(p));await record('ABOVE-200-HISTORY',p,lambda:above(p,context));cleanup()
            await record('KEYBOARD-DOM',p,lambda:keyboard(p));await record('SELECTED-XLSX',p,lambda:download(p));await context.close();await browser.close()
    finally:cleanup()
    assert initial==db.rows('SELECT COUNT(*) AS n FROM dbo.m_employee')[0]['n'];assert artifact==h.func.artifact()
    result={'runId':'sp2-ui-regenerate-20261029','recordedAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'artifactSha256':artifact,'testSha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'database':db.DATABASE,'results':results,'observed':observed,'pageErrors':errors,'boundaryRowsAfter':0,'formalScore':None};(OUT/'result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf8');return int(any(r['status']=='FAIL' for r in results) or bool(errors))
if __name__=='__main__':sys.exit(asyncio.run(main()))
